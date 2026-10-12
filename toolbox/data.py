from __future__ import annotations

import csv
import io
import json
import math
import sqlite3
import statistics
import time

from .core import CONFIG, FS_LOCK, LIMIT, STATE, TEXT_LIMIT, backup, bounded, input_file, resolve, tool, write_bytes


@tool("data")
def json_read(path: str) -> dict:
    """Parse a JSON file; refuse values too large for a tool response."""
    p = input_file(path)
    text = p.read_text(encoding="utf-8-sig")
    if len(text) > TEXT_LIMIT:
        raise ValueError("JSON exceeds response limit")
    return {"value": json.loads(text)}


@tool("data", write=True)
def json_write(path: str, value: dict | list, overwrite: bool = False) -> dict:
    """Write a JSON object or array as formatted UTF-8."""
    return write_bytes(path, json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8"), overwrite)


def csv_rows(path: str, delimiter: str = ","):
    if len(delimiter) != 1:
        raise ValueError("delimiter must be one character")
    p = input_file(path)
    return csv.DictReader(io.StringIO(p.read_text(encoding="utf-8-sig")), delimiter=delimiter)


@tool("data")
def csv_read(path: str, delimiter: str = ",", offset: int = 0, limit: int = 100) -> dict:
    """Read a page of CSV rows with a header; supports UTF-8 BOM files."""
    bounded(offset, low=0, high=1_000_000)
    bounded(limit)
    reader = csv_rows(path, delimiter)
    rows, chars = [], 0
    for index, row in enumerate(reader):
        if index < offset:
            continue
        size = len(json.dumps(row, ensure_ascii=False))
        if len(rows) >= limit or chars + size > TEXT_LIMIT:
            return {"columns": reader.fieldnames, "rows": rows, "truncated": True}
        rows.append(row)
        chars += size
    return {"columns": reader.fieldnames, "rows": rows, "truncated": False}


@tool("data", write=True)
def csv_write(path: str, columns: list[str], rows: list[list], delimiter: str = ",",
              overwrite: bool = False) -> dict:
    """Write tabular data to a UTF-8 BOM CSV; all rows must match column count."""
    if len(delimiter) != 1 or not columns or len(rows) > 10000:
        raise ValueError("Invalid table or delimiter")
    if any(len(row) != len(columns) for row in rows):
        raise ValueError("Row width must match columns")
    out = io.StringIO(newline="")
    writer = csv.writer(out, delimiter=delimiter)
    writer.writerow(columns)
    writer.writerows(rows)
    return write_bytes(path, out.getvalue().encode("utf-8-sig"), overwrite)


@tool("data")
def csv_summary(path: str, delimiter: str = ",") -> dict:
    """Summarize up to 100,000 CSV rows: missing values and finite numeric statistics."""
    reader = csv_rows(path, delimiter)
    if not reader.fieldnames or len(reader.fieldnames) > 200:
        raise ValueError("Expected at most 200 named columns")
    numbers = {key: [] for key in reader.fieldnames}
    missing = {key: 0 for key in reader.fieldnames}
    count = 0
    for row in reader:
        count += 1
        if count > 100000:
            raise ValueError("CSV exceeds 100,000 rows")
        for key in reader.fieldnames:
            value = row.get(key)
            if value is None or not value.strip():
                missing[key] += 1
                continue
            try:
                number = float(value)
                if math.isfinite(number):
                    numbers[key].append(number)
            except (ValueError, TypeError):
                pass
    summary = {}
    for key, values in numbers.items():
        summary[key] = {"missing": missing[key], "numeric_count": len(values)}
        if values:
            summary[key].update(min=min(values), max=max(values), mean=statistics.mean(values),
                                median=statistics.median(values), sum=sum(values))
    return {"row_count": count, "columns": summary}


@tool("data")
def xlsx_read(path: str, sheet: str = "", start_row: int = 1, row_count: int = 100,
              values_only: bool = False) -> dict:
    """Read an XLSX sheet page. values_only reads cached formula values without recalculating."""
    from openpyxl import load_workbook
    bounded(start_row, high=1_048_576)
    bounded(row_count)
    wb = load_workbook(input_file(path), read_only=True, data_only=values_only)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows, chars = [], 0
        truncated = False
        for row in ws.iter_rows(min_row=start_row, max_row=min(start_row + row_count - 1, ws.max_row or 1),
                                max_col=min(ws.max_column or 1, 200), values_only=True):
            row = [v.isoformat() if hasattr(v, "isoformat") else v for v in row]
            size = len(json.dumps(row, default=str, ensure_ascii=False))
            if chars + size > TEXT_LIMIT:
                truncated = True
                break
            rows.append(row)
            chars += size
        return {"sheets": wb.sheetnames, "sheet": ws.title, "total_rows": ws.max_row,
                "total_columns": ws.max_column, "rows": rows,
                "truncated": truncated or (ws.max_row or 0) >= start_row + row_count or (ws.max_column or 0) > 200,
                "formula_note": "Cached formula values may be absent or stale; formulas are not recalculated."}
    finally:
        wb.close()


@tool("data", write=True)
def xlsx_write(path: str, columns: list[str], rows: list[list], sheet: str = "Data",
               overwrite: bool = False) -> dict:
    """Create a styled XLSX table with autofilter and frozen headers; input strings are literal."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if not columns or len(columns) > 200 or len(rows) > 10000:
        raise ValueError("Table must have 1..200 columns and at most 10,000 rows")
    if any(len(row) != len(columns) for row in rows):
        raise ValueError("Row width must match columns")
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(columns)
    for row in rows:
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245B78")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(14, len(col) + 4))
    out = io.BytesIO()
    wb.save(out)
    return write_bytes(path, out.getvalue(), overwrite)


def connection(path: str, write: bool = False):
    p = resolve(path) if write else input_file(path)
    con = sqlite3.connect(p if write else p.as_uri() + "?mode=ro", uri=not write, timeout=5)
    con.row_factory = sqlite3.Row
    con.enable_load_extension(False)
    denied = {sqlite3.SQLITE_ATTACH, sqlite3.SQLITE_DETACH, sqlite3.SQLITE_PRAGMA}
    def authorize(action, a, b, db, trigger):
        if action in denied or (action == sqlite3.SQLITE_FUNCTION and (b or "").lower() == "load_extension"):
            return sqlite3.SQLITE_DENY
        return sqlite3.SQLITE_OK
    con.set_authorizer(authorize)
    deadline = time.monotonic() + 5
    con.set_progress_handler(lambda: int(time.monotonic() > deadline), 1000)
    return con


@tool("sqlite")
def sqlite_schema(path: str) -> dict:
    """Read SQLite table, view and index definitions without modifying the database."""
    con = connection(path)
    try:
        rows = con.execute("SELECT type,name,tbl_name,sql FROM sqlite_master WHERE name NOT LIKE 'sqlite_%' ORDER BY type,name LIMIT 201").fetchall()
        return {"objects": [dict(row) for row in rows[:200]], "truncated": len(rows) > 200}
    finally:
        con.close()


@tool("sqlite")
def sqlite_query(path: str, sql: str, parameters: list | None = None, limit: int = 200) -> dict:
    """Run one parameterized read-only SQLite statement with row, response and time limits."""
    bounded(limit)
    con = connection(path)
    try:
        cursor = con.execute(sql, parameters or [])
        if cursor.description is None:
            raise ValueError("Expected a query result")
        rows = []
        chars = 0
        truncated = False
        for row in cursor.fetchmany(limit + 1):
            item = {key: ({"blob_hex": value.hex()} if isinstance(value, bytes) else value)
                    for key, value in dict(row).items()}
            size = len(json.dumps(item, ensure_ascii=False))
            if len(rows) >= limit or chars + size > TEXT_LIMIT:
                truncated = True
                break
            rows.append(item)
            chars += size
        return {"columns": [column[0] for column in cursor.description], "rows": rows, "truncated": truncated}
    finally:
        con.close()


@tool("sqlite", write=True, destructive=True)
def sqlite_execute(path: str, sql: str, parameters: list | None = None) -> dict:
    """Run one SQLite write statement if sqlite_writes is enabled; back up first and commit atomically."""
    if not CONFIG.get("sqlite_writes", False):
        raise ValueError("Set sqlite_writes=true in settings.json to enable DB writes")
    with FS_LOCK:
        p = resolve(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        # SQLite online backup includes committed WAL content.
        saved = None
        if p.exists():
            saved = backup(input_file(path))
            original = sqlite3.connect(p)
            target = sqlite3.connect(str(STATE / "backups" / saved))
            try:
                original.backup(target)
            finally:
                target.close()
                original.close()
        con = connection(path, True)
        try:
            with con:
                cursor = con.execute(sql, parameters or [])
                if cursor.description is not None:
                    raise ValueError("Use sqlite_query for queries and RETURNING statements")
            return {"path": str(p), "affected_rows": cursor.rowcount, "backup_id": saved}
        finally:
            con.close()
