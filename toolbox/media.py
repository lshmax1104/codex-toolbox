from __future__ import annotations

import io

from .core import LIMIT, TEXT_LIMIT, bounded, input_file, tool, write_bytes


@tool("documents")
def document_extract(path: str, start_page: int = 1, page_count: int = 20) -> dict:
    """Extract PDF text, DOCX paragraphs/tables or plain UTF-8. Scanned PDFs need separate OCR."""
    p = input_file(path)
    bounded(start_page, high=100000)
    bounded(page_count, high=100)
    if p.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(p)
        pages, chars = [], 0
        truncated = False
        for i in range(start_page - 1, min(len(reader.pages), start_page - 1 + page_count)):
            value = reader.pages[i].extract_text() or ""
            text = value[:max(0, TEXT_LIMIT - chars)]
            pages.append({"page": i + 1, "text": text})
            chars += len(text)
            if len(text) < len(value) or chars >= TEXT_LIMIT:
                truncated = True
                break
        return {"path": str(p), "total_pages": len(reader.pages), "pages": pages,
                "truncated": truncated or len(reader.pages) >= start_page + page_count}
    if p.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(p)
        pieces = [paragraph.text for paragraph in doc.paragraphs]
        for table in doc.tables:
            pieces.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(pieces)
    else:
        text = p.read_text(encoding="utf-8-sig")
    return {"path": str(p), "text": text[:TEXT_LIMIT], "truncated": len(text) > TEXT_LIMIT}


@tool("documents", write=True)
def pdf_merge(paths: list[str], destination: str, overwrite: bool = False) -> dict:
    """Merge up to 50 PDFs in input order, capped at 1,000 pages and max_file_bytes."""
    from pypdf import PdfReader, PdfWriter
    bounded(len(paths), high=50)
    writer = PdfWriter()
    count = 0
    for path in paths:
        reader = PdfReader(input_file(path))
        count += len(reader.pages)
        if count > 1000:
            raise ValueError("Combined PDFs exceed 1,000 pages")
        writer.append(reader)
    out = io.BytesIO()
    writer.write(out)
    return {**write_bytes(destination, out.getvalue(), overwrite), "pages": count}


@tool("documents", write=True)
def pdf_select_pages(path: str, pages: list[int], destination: str, overwrite: bool = False) -> dict:
    """Extract or reorder PDF pages using explicit one-based page numbers."""
    from pypdf import PdfReader, PdfWriter
    bounded(len(pages), high=1000)
    reader, writer = PdfReader(input_file(path)), PdfWriter()
    for number in pages:
        bounded(number, high=len(reader.pages))
        writer.add_page(reader.pages[number - 1])
    out = io.BytesIO()
    writer.write(out)
    return {**write_bytes(destination, out.getvalue(), overwrite), "pages": len(pages)}


@tool("images")
def image_info(path: str) -> dict:
    """Inspect image format, dimensions, color mode and frame count."""
    from PIL import Image
    with Image.open(input_file(path)) as im:
        return {"format": im.format, "width": im.width, "height": im.height,
                "mode": im.mode, "frames": getattr(im, "n_frames", 1)}


def image_output(image, destination, format, overwrite):
    if format.upper() not in {"PNG", "JPEG", "WEBP"}:
        raise ValueError("format must be PNG, JPEG or WEBP")
    if format.upper() == "JPEG":
        image = image.convert("RGB")
    out = io.BytesIO()
    image.save(out, format=format.upper())
    return write_bytes(destination, out.getvalue(), overwrite)


@tool("images", write=True)
def image_resize(path: str, destination: str, width: int, height: int,
                  keep_aspect: bool = True, format: str = "PNG", overwrite: bool = False) -> dict:
    """Resize an image with orientation correction; keep_aspect fits it inside the requested box."""
    from PIL import Image, ImageOps
    bounded(width, high=8192)
    bounded(height, high=8192)
    with Image.open(input_file(path)) as source:
        im = ImageOps.exif_transpose(source)
        if keep_aspect:
            im.thumbnail((width, height), Image.Resampling.LANCZOS)
        else:
            im = im.resize((width, height), Image.Resampling.LANCZOS)
        return {**image_output(im, destination, format, overwrite), "width": im.width, "height": im.height}


@tool("images", write=True)
def image_convert(path: str, destination: str, format: str = "PNG", overwrite: bool = False) -> dict:
    """Convert the first image frame to PNG, JPEG or WEBP with orientation correction."""
    from PIL import Image, ImageOps
    with Image.open(input_file(path)) as source:
        return image_output(ImageOps.exif_transpose(source), destination, format, overwrite)
