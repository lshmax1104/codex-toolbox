"""Protocol-level acceptance suite. All mutations use a new isolated test folder."""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
import uuid
import zipfile
from datetime import timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

BASE = Path(__file__).resolve().parent


class FixtureHTTP(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == "/redirect":
            self.send_response(302)
            self.send_header("Location", "/page")
            self.end_headers()
            return
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.end_headers()
        self.wfile.write(b"<html><title>Fixture</title><script>SECRET_SCRIPT</script><p>Hello MCP</p></html>")

    def do_POST(self):
        size = int(self.headers.get("Content-Length", "0"))
        body = self.rfile.read(size)
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"body": json.loads(body), "authorization": self.headers.get("Authorization")}).encode())

    def log_message(self, format, *args):
        pass


async def suite(root: Path):
    from PIL import Image
    from docx import Document
    from pypdf import PdfWriter

    root.mkdir(parents=True)
    http = ThreadingHTTPServer(("127.0.0.1", 0), FixtureHTTP)
    thread = threading.Thread(target=http.serve_forever, daemon=True)
    thread.start()
    url = f"http://127.0.0.1:{http.server_port}"
    settings_path = BASE / "settings.json"
    if not settings_path.exists():
        settings_path = BASE / "settings.example.json"
    settings = json.loads(settings_path.read_text(encoding="utf-8-sig"))
    settings.update(allowed_roots=[str(root)], state_dir=str(root / "internal-state"),
                    sqlite_writes=True, web_allow_private_network=True)
    settings["services"] = {"fixture": {"base_url": url, "methods": ["GET", "POST"],
        "allow_private_network": True, "token_env": "TOOLBOX_TEST_TOKEN"}}
    settings["tasks"] = {
        "hello": {"command": ["{python}", "-c", "print('TASK_OK')"], "timeout_seconds": 10},
        "slow": {"command": ["{python}", "-c", "import time; time.sleep(20)"], "timeout_seconds": 30},
        "timeout": {"command": ["{python}", "-c", "import time; time.sleep(20)"], "timeout_seconds": 1},
        "large": {"command": ["{python}", "-c", "print('x'*300000)"], "timeout_seconds": 10},
    }
    config = root / "test-settings.json"
    config.write_text(json.dumps(settings), encoding="utf-8")
    (root / "original.txt").write_text("안녕하세요\nalpha beta\n", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("Document fixture")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table fixture"
    doc.save(root / "fixture.docx")
    pdf = PdfWriter()
    pdf.add_blank_page(300, 400)
    pdf.add_blank_page(300, 400)
    with (root / "fixture.pdf").open("wb") as f:
        pdf.write(f)
    Image.new("RGBA", (120, 80), (20, 40, 80, 128)).save(root / "fixture.png")
    with sqlite3.connect(root / "fixture.db") as con:
        con.execute("CREATE TABLE sales(name TEXT, amount INTEGER)")
        con.executemany("INSERT INTO sales VALUES (?,?)", [("a", 10), ("b", 20)])
    git = settings.get("git_executable") or shutil.which("git")
    repo = root / "repo"
    repo.mkdir()
    def git_run(*args):
        subprocess.run([git, "-C", str(repo), *args], check=True, capture_output=True,
                       creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    git_run("init")
    (repo / "hello.txt").write_text("before\n", encoding="utf-8")
    git_run("add", "hello.txt")
    git_run("-c", "user.name=Toolbox Test", "-c", "user.email=toolbox@example.invalid", "commit", "-m", "fixture")
    (repo / "hello.txt").write_text("after\n", encoding="utf-8")
    checks, called = [], set()
    params = StdioServerParameters(command=sys.executable, args=[str(BASE / "server.py")],
        env=dict(os.environ, TOOLBOX_CONFIG=str(config), TOOLBOX_TEST_TOKEN="fixture-token-secret"))
    try:
        async with stdio_client(params) as (read, write):
            async with ClientSession(read, write, read_timeout_seconds=timedelta(seconds=45)) as client:
                init = await client.initialize()
                assert init.serverInfo.name == "codex-toolbox"
                tools = (await client.list_tools()).tools
                assert len(tools) == 57
                assert len({t.name for t in tools}) == len(tools)
                for t in tools:
                    assert t.description and t.inputSchema["type"] == "object"
                    assert t.annotations is not None
                    assert t.outputSchema and t.outputSchema["type"] == "object"
                checks.append("MCP initialize, discovery and all 57 tool schemas")

                async def call(name, args=None, error=False):
                    called.add(name)
                    result = await client.call_tool(name, args or {})
                    if error:
                        assert result.isError, f"{name}: expected error"
                        return result
                    assert not result.isError, f"{name}: {result.content}"
                    assert result.structuredContent is not None, f"{name}: missing structured output"
                    if result.structuredContent is not None:
                        return result.structuredContent
                    return json.loads(result.content[0].text)

                health = await call("toolbox_health")
                assert health["tool_count"] == 57 and health["allowed_roots"] == [str(root)]
                assert (await call("system_info"))["memory_total"] > 0
                await call("processes_list", {"limit": 5})
                await call("ports_list", {"limit": 5})
                assert "안녕하세요" in (await call("files_read", {"path": "original.txt"}))["text"]
                assert len((await call("files_search", {"path": ".", "query": "alpha", "pattern": "*.txt"}))["matches"]) == 1
                await call("files_list", {"recursive": True, "pattern": "*.txt"})
                await call("files_mkdir", {"path": "created"})
                await call("files_write", {"path": "created/new.txt", "text": "one"})
                await call("files_write", {"path": "created/new.txt", "text": "bad"}, error=True)
                changed = await call("files_replace", {"path": "created/new.txt", "old": "one", "new": "two"})
                await call("files_replace", {"path": "created/new.txt", "old": "absent", "new": "bad"}, error=True)
                await call("backups_restore", {"backup_id": changed["backup_id"], "destination": "restored.txt"})
                assert (root / "restored.txt").read_text() == "one"
                assert len((await call("backups_list"))["backups"]) >= 1
                await call("files_copy", {"source": "original.txt", "destination": "copied.txt"})
                assert (await call("files_compare", {"first": "original.txt", "second": "copied.txt"}))["identical"]
                assert len((await call("files_hash", {"path": "copied.txt"}))["sha256"]) == 64
                assert (await call("files_info", {"path": "original.txt"}))["bytes"] > 0
                assert len((await call("files_duplicates", {"path": ".", "pattern": "*.txt"}))["groups"]) >= 1
                await call("files_move", {"source": "copied.txt", "destination": "moved.txt"})
                assert not (root / "copied.txt").exists()
                await call("files_read", {"path": str(root.parent / "outside.txt")}, error=True)
                await call("files_write", {"path": str(config), "text": "{}", "overwrite": True}, error=True)
                await call("files_write", {"path": "internal-state/bad.txt", "text": "bad"}, error=True)
                checks.append("UTF-8 files, exact replacement, backup/restore, duplicates and path restrictions")

                await call("zip_create", {"paths": ["original.txt", "moved.txt"], "destination": "bundle.zip"})
                assert len((await call("zip_list", {"path": "bundle.zip"}))["entries"]) == 2
                await call("zip_extract", {"path": "bundle.zip", "destination": "unpacked"})
                assert (root / "unpacked/original.txt").read_bytes() == (root / "original.txt").read_bytes()
                with zipfile.ZipFile(root / "evil.zip", "w") as z:
                    z.writestr("../escaped.txt", "bad")
                await call("zip_extract", {"path": "evil.zip", "destination": "rejected"}, error=True)
                assert not (root / "rejected").exists()
                checks.append("ZIP round trip and traversal rejection")

                await call("json_write", {"path": "data.json", "value": {"한국어": [1, 2]}})
                assert (await call("json_read", {"path": "data.json"}))["value"]["한국어"] == [1, 2]
                await call("csv_write", {"path": "data.csv", "columns": ["name", "amount"], "rows": [["가", 10], ["나", 20], ["다", ""]]})
                assert len((await call("csv_read", {"path": "data.csv", "offset": 1}))["rows"]) == 2
                summary = await call("csv_summary", {"path": "data.csv"})
                assert summary["columns"]["amount"]["sum"] == 30 and summary["columns"]["amount"]["missing"] == 1
                await call("xlsx_write", {"path": "data.xlsx", "columns": ["name", "amount"], "rows": [["가", 10], ["=1+1", 20]]})
                sheet = await call("xlsx_read", {"path": "data.xlsx"})
                assert sheet["rows"][1] == ["가", 10] and sheet["rows"][2][0] == "=1+1"
                checks.append("JSON/CSV/XLSX round trips, numeric summary and literal spreadsheet cells")

                doc = await call("document_extract", {"path": "fixture.docx"})
                assert "Document fixture" in doc["text"] and "Table fixture" in doc["text"]
                assert (await call("document_extract", {"path": "fixture.pdf"}))["total_pages"] == 2
                await call("pdf_merge", {"paths": ["fixture.pdf", "fixture.pdf"], "destination": "merged.pdf"})
                await call("pdf_select_pages", {"path": "merged.pdf", "pages": [4, 1], "destination": "selected.pdf"})
                assert (await call("document_extract", {"path": "selected.pdf"}))["total_pages"] == 2
                assert (await call("image_info", {"path": "fixture.png"}))["width"] == 120
                await call("image_resize", {"path": "fixture.png", "destination": "small.png", "width": 60, "height": 60})
                assert (await call("image_info", {"path": "small.png"}))["height"] == 40
                await call("image_convert", {"path": "fixture.png", "destination": "converted.jpg", "format": "JPEG"})
                assert (await call("image_info", {"path": "converted.jpg"}))["mode"] == "RGB"
                checks.append("DOCX extraction, PDF merge/page selection and image resizing/conversion")

                assert (await call("sqlite_schema", {"path": "fixture.db"}))["objects"][0]["name"] == "sales"
                result = await call("sqlite_query", {"path": "fixture.db", "sql": "SELECT sum(amount) total FROM sales WHERE amount>=?", "parameters": [10]})
                assert result["rows"][0]["total"] == 30
                await call("sqlite_query", {"path": "fixture.db", "sql": "DELETE FROM sales"}, error=True)
                await call("sqlite_query", {"path": "fixture.db", "sql": "ATTACH DATABASE 'other.db' AS other"}, error=True)
                result = await call("sqlite_execute", {"path": "fixture.db", "sql": "INSERT INTO sales VALUES (?,?)", "parameters": ["c", 30]})
                assert result["affected_rows"] == 1 and result["backup_id"]
                await call("backups_restore", {"backup_id": result["backup_id"], "destination": "db-restored.db"})
                assert (await call("sqlite_query", {"path": "db-restored.db", "sql": "SELECT count(*) total FROM sales"}))["rows"][0]["total"] == 2
                checks.append("Parameterized SQLite queries, read-only enforcement, DB write and online backup")

                assert "hello.txt" in (await call("git_status", {"path": "repo"}))["output"]
                assert "+after" in (await call("git_diff", {"path": "repo"}))["output"]
                assert "fixture" in (await call("git_log", {"path": "repo"}))["output"]
                assert "fixture" in (await call("git_show", {"path": "repo"}))["output"]
                assert (await call("git_branches", {"path": "repo"}))["output"].strip()
                await call("git_show", {"path": "repo", "ref": "--help"}, error=True)
                checks.append("Git status/diff/log/show/branches and argument validation")

                page = await call("web_fetch", {"url": url + "/redirect"})
                assert "Hello MCP" in page["text"] and "SECRET_SCRIPT" not in page["text"]
                await call("web_download", {"url": url + "/page", "destination": "download.html"})
                assert (root / "download.html").exists()
                assert (await call("services_list"))["services"][0]["configured"]
                response = await call("service_request", {"service": "fixture", "path": "echo", "method": "POST", "body": {"test": 1}})
                assert "fixture-token-secret" not in response["text"] and "[REDACTED]" in response["text"]
                await call("service_request", {"service": "fixture", "path": "https://example.com"}, error=True)
                await call("service_request", {"service": "fixture", "method": "DELETE"}, error=True)
                await call("web_fetch", {"url": "file:///C:/Windows/win.ini"}, error=True)
                checks.append("HTTP redirect/text/download, configured REST methods and credential redaction")

                await call("memory_put", {"namespace": "test", "key": "한국어", "value": "persistent note"})
                assert (await call("memory_get", {"namespace": "test", "key": "한국어"}))["value"] == "persistent note"
                assert len((await call("memory_search", {"query": "persistent"}))["memories"]) == 1
                await call("memory_delete", {"namespace": "test", "key": "한국어"})
                assert not (await call("memory_get", {"namespace": "test", "key": "한국어"}))["found"]
                await call("memory_put", {"namespace": "test", "key": "restart", "value": "survives restart"})
                checks.append("Namespaced memory create/read/search/delete")

                assert len((await call("tasks_list"))["tasks"]) == 4
                assert "TASK_OK" in (await call("task_run", {"name": "hello"}))["output"]
                large = await call("task_run", {"name": "large"})
                assert large["truncated"] and len(large["output"]) == 60000
                timeout = await call("task_run", {"name": "timeout", "wait_seconds": 5})
                assert timeout["status"] == "timed_out"
                job = await call("job_start", {"name": "slow"})
                assert (await call("job_status", {"job_id": job["id"]}))["status"] == "running"
                cancelled = await call("job_cancel", {"job_id": job["id"]})
                assert cancelled["status"] == "cancelled"
                assert (await call("jobs_list"))["jobs"]
                await call("task_run", {"name": "arbitrary-shell"}, error=True)
                assert (await call("audit_tail"))["events"]
                checks.append("Named task execution, bounded output, timeout, cancellation and audit metadata")

                resources = (await client.list_resources()).resources
                assert len(resources) == 2
                content = await client.read_resource("toolbox://capabilities")
                assert len(json.loads(content.contents[0].text)) == 57
                prompts = (await client.list_prompts()).prompts
                assert len(prompts) == 3
                prompt = await client.get_prompt("inspect_project", {"path": "repo"})
                assert prompt.messages
                assert called == {tool.name for tool in tools}, "Missing calls: " + str({tool.name for tool in tools} - called)
                checks.append("Resources, prompts and every discovered tool called through stdio MCP")
        # A new server process must read the existing memory state.
        async with stdio_client(params) as (read, write):
            async with ClientSession(read, write, read_timeout_seconds=timedelta(seconds=45)) as client:
                await client.initialize()
                result = await client.call_tool("memory_get", {"namespace": "test", "key": "restart"})
                assert not result.isError, str(result.content)
                value = result.structuredContent or json.loads(result.content[0].text)
                assert value["value"] == "survives restart"
                checks.append("Memory persistence across server restart")
    finally:
        http.shutdown()
        http.server_close()
    return {"passed": True, "tool_count": len(called), "checks": checks, "test_root": str(root)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workdir", type=Path, default=BASE.parent.parent / "work" / "mcp-build")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    root = args.workdir.resolve() / ("acceptance-" + uuid.uuid4().hex[:12])
    result = asyncio.run(suite(root))
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.write_text(text, encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
